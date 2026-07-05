from io import BytesIO

from openpyxl import load_workbook

from app import create_app
from app.config import Config
from app.services.database_service import DatabaseService


class ReportsApiConfig(Config):
    TESTING = True
    SECRET_KEY = "test-secret"
    ADMIN_USERNAME = "admin"
    ADMIN_PASSWORD = "admin123"
    DATABASE_URL = "postgresql://localhost/smartbi"


def _login(client) -> None:
    login_response = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    assert login_response.status_code == 200


def _patch_uploads(monkeypatch, datasets):
    def fetch_upload_dataset(self, upload_id=None):
        item = datasets.get(upload_id)
        if item is None:
            return None, []
        return item["upload"], item["rows"]

    monkeypatch.setattr(DatabaseService, "fetch_upload_dataset", fetch_upload_dataset)


def _upload(upload_id, file_name, domain_name, rows):
    return {
        "upload": {
            "id": upload_id,
            "file_name": file_name,
            "domain_name": domain_name,
            "confidence": 0.91,
            "row_count": len(rows),
            "summary_json": {},
            "created_at": "2026-01-01T00:00:00",
        },
        "rows": rows,
    }


def _sales_rows():
    return [
        {"Order ID": "A1", "Order Date": "2026-01-01", "Customer": "Nia", "Product": "Laptop", "Region": "West", "Revenue Amount": 1000},
        {"Order ID": "A2", "Order Date": "2026-01-02", "Customer": "Dev", "Product": "Mouse", "Region": "East", "Revenue Amount": 500},
    ]


def test_reports_metadata_endpoint_returns_available_formats() -> None:
    app = create_app(ReportsApiConfig)
    with app.test_client() as client:
        _login(client)
        response = client.get("/api/reports")

    assert response.status_code == 200
    payload = response.get_json()
    assert payload is not None
    assert "pdf" in payload["available_formats"]
    assert "excel" in payload["available_formats"]
    assert payload["download_endpoint"] == "/api/reports/download"


def test_pdf_report_download(monkeypatch) -> None:
    _patch_uploads(monkeypatch, {1: _upload(1, "sales.xlsx", "Retail Sales", _sales_rows())})

    app = create_app(ReportsApiConfig)
    with app.test_client() as client:
        _login(client)
        response = client.get("/api/reports/download?upload_id=1&format=pdf")

    assert response.status_code == 200
    assert response.mimetype == "application/pdf"
    assert response.data.startswith(b"%PDF")


def test_excel_report_download(monkeypatch) -> None:
    _patch_uploads(monkeypatch, {1: _upload(1, "sales.xlsx", "Retail Sales", _sales_rows())})

    app = create_app(ReportsApiConfig)
    with app.test_client() as client:
        _login(client)
        response = client.get("/api/reports/download?upload_id=1&format=excel")

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data))
    assert {"Summary", "Analytics", "Raw Data"}.issubset(set(workbook.sheetnames))
    assert workbook["Summary"]["A1"].value == "SmartBI Report"


def test_csv_report_download(monkeypatch) -> None:
    _patch_uploads(monkeypatch, {1: _upload(1, "sales.xlsx", "Retail Sales", _sales_rows())})

    app = create_app(ReportsApiConfig)
    with app.test_client() as client:
        _login(client)
        response = client.get("/api/reports/download?upload_id=1&format=csv")

    assert response.status_code == 200
    assert response.mimetype == "text/csv"
    assert "SmartBI Report" in response.data.decode("utf-8")
    assert "Total Revenue" in response.data.decode("utf-8")


def test_comparison_report_download(monkeypatch) -> None:
    old_rows = _sales_rows()
    new_rows = [
        {"Order ID": "B1", "Order Date": "2026-02-01", "Customer": "Nia", "Product": "Laptop", "Region": "West", "Revenue Amount": 1200},
        {"Order ID": "B2", "Order Date": "2026-02-02", "Customer": "Mira", "Product": "Keyboard", "Region": "South", "Revenue Amount": 900},
    ]
    _patch_uploads(monkeypatch, {
        1: _upload(1, "sales_old.xlsx", "Retail Sales", old_rows),
        2: _upload(2, "sales_new.xlsx", "Retail Sales", new_rows),
    })

    app = create_app(ReportsApiConfig)
    with app.test_client() as client:
        _login(client)
        response = client.get("/api/reports/download?upload_id=1&format=csv&compare=true&upload_b=2")

    assert response.status_code == 200
    text = response.data.decode("utf-8")
    assert "Comparison" not in text
    assert "Growth %" in text
    assert "Total Revenue" in text


def test_report_download_rejects_invalid_format(monkeypatch) -> None:
    _patch_uploads(monkeypatch, {1: _upload(1, "sales.xlsx", "Retail Sales", _sales_rows())})

    app = create_app(ReportsApiConfig)
    with app.test_client() as client:
        _login(client)
        response = client.get("/api/reports/download?upload_id=1&format=xml")

    assert response.status_code == 400
    payload = response.get_json()
    assert payload is not None
    assert payload["message"] == "Unsupported report format."


def test_report_download_returns_404_for_missing_upload(monkeypatch) -> None:
    _patch_uploads(monkeypatch, {})

    app = create_app(ReportsApiConfig)
    with app.test_client() as client:
        _login(client)
        response = client.get("/api/reports/download?upload_id=99&format=pdf")

    assert response.status_code == 404
    payload = response.get_json()
    assert payload is not None
    assert payload["message"] == "Upload does not exist."


def test_report_download_rejects_empty_dataset(monkeypatch) -> None:
    _patch_uploads(monkeypatch, {1: _upload(1, "empty.xlsx", "Retail Sales", [])})

    app = create_app(ReportsApiConfig)
    with app.test_client() as client:
        _login(client)
        response = client.get("/api/reports/download?upload_id=1&format=csv")

    assert response.status_code == 400
    payload = response.get_json()
    assert payload is not None
    assert payload["message"] == "Upload dataset is empty."


def test_report_download_requires_comparison_upload(monkeypatch) -> None:
    _patch_uploads(monkeypatch, {1: _upload(1, "sales.xlsx", "Retail Sales", _sales_rows())})

    app = create_app(ReportsApiConfig)
    with app.test_client() as client:
        _login(client)
        response = client.get("/api/reports/download?upload_id=1&format=csv&compare=true")

    assert response.status_code == 400
    payload = response.get_json()
    assert payload is not None
    assert payload["message"] == "upload_b is required when compare=true."
